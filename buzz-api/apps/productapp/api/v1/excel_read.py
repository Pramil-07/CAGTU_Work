import os
import uuid
from itertools import islice
from openpyxl import load_workbook
from openpyxl_image_loader import SheetImageLoader
from apps.productapp.models import FileStore
from django.conf import settings



"""
Inside openpyxl_image_loader venv package; customize col; package; SheetImageLoader (customize it and replace the col)
col = openpyxl.utils.cell.get_column_letter(image.anchor._from.col + 1)
"""
class Workbook:

    def __init__(self, excelfile,request):
        # load the workbook
        self.wb = load_workbook(excelfile)
        self.mobile_sheet = self.wb["Mobiles"]
        self.request = request
        for row in self.mobile_sheet.values:
            if row[0] == None:
                break
            product_number = row[0]
        self.product_number = int(product_number)

    def image(self):
        """
        Image function to grab image from excel sheet and save it in filestore
        """
        image_urls = []
        image_loader = SheetImageLoader(self.mobile_sheet)
        
        try:
            for count in range(3, self.product_number + 3):
                image_cell = "B"+str(count)
                image = image_loader.get(image_cell)
                filename = str(uuid.uuid4())
                image_path = os.path.join(settings.MEDIA_ROOT, "filestore")
                image.save(f'{image_path}/{filename}.png')

                try:
                    image_data = f"filestore/{filename}.png"
                    create = FileStore.objects.create(image=image_data)
                    image_link = self.request.build_absolute_uri(create.image.url)
                    image_urls.append(image_link)
                    
                except Exception as e:
                    print("Exception", e)
        except Exception as e:
            print("Exception", e)
        return image_urls


    def stock_image(self):
        """
        Image function to grab image from excel sheet and save it in filestore
        """
        image_urls = []
        image_loader = SheetImageLoader(self.mobile_sheet)
        
        for count in range(3, self.mobile_sheet.max_row + 1):
            image_cell = "AF"+str(count)
            image = image_loader.get(image_cell)
            filename = str(uuid.uuid4())
            image_path = os.path.join(settings.MEDIA_ROOT, "filestore")
            image.save(f'{image_path}/{filename}.png')

            try:
                image_data = f"filestore/{filename}.png"
                create = FileStore.objects.create(image=image_data)
                image_link = self.request.build_absolute_uri(create.image.url)
                image_urls.append(image_link)
                
            except Exception as e:
                print("Exception", e)
            
        return image_urls

    
    def create_product(self):
        """
        for creating a product list from the product
        """
        product_attribute_list = []

        try:
            for row in islice(self.mobile_sheet.values,2,self.product_number+2):
                attribute_list = {}
                attribute_list["Product Number"] = row[0]
                attribute_list["Category"] = row[2]
                attribute_list["Is_Active"] = row[3]
                attribute_list["Name"] = row[4]
                attribute_list["Model_no"] = row[5]
                attribute_list["Video_url"] = row[6]
                attribute_list["Product_status"] = row[7]
                attribute_list["Description"] = row[8]
                attribute_list["Warranty_type"] = row[9]
                attribute_list["Warranty_period"] = row[10]
                attribute_list["Notes"] = row[11]
                attribute_list["Meta_title"] = row[12]
                attribute_list["Meta_description"] = row[13]
                attribute_list["Meta_keyword"] = row[14]
                attribute_list["Brand"] = row[15]
                attribute_list["Screen_size"] = row[16]
                attribute_list["Screen_type"] = row[17]
                attribute_list["Processor"] = row[18]
                attribute_list["Camera_front"] = row[19]
                attribute_list["Camera_back"] = row[20]
                product_attribute_list.append(attribute_list)
            return product_attribute_list
        except Exception as e:
            print("Exception", e)

    def create_stocks(self):
        stock_attribute_list = []

        for row in islice(self.mobile_sheet.values,2,self.mobile_sheet.max_row):
            attribute_list = {}
            attribute_list["Product_code"] = row[21]
            attribute_list["RAM"] = row[22]
            attribute_list["ROM"] = row[23]
            attribute_list["MRP"] = row[24]
            attribute_list["Price"] = row[25]
            attribute_list["Color"] = row[26]
            attribute_list["Availability"] = row[27]
            attribute_list["Size_unit"] = row[28]
            attribute_list["Size"] = row[29]
            attribute_list["Quantity"] = row[30]
            # attribute_list["Image"] = row[30]
            attribute_list["Is_default"] = row[32]
            stock_attribute_list.append(attribute_list)
        return stock_attribute_list